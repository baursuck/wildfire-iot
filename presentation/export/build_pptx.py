"""Build a .pptx from the HTML deck: each slide is a 1920x1080 still rendered with ?static,
slide 6 gets the detection video as a native, auto-playing, looping movie."""
import copy
import sys

from lxml import etree
from pptx import Presentation
from pptx.util import Emu, Inches

renders, video, poster, out = sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4]
NS = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
P = "{%s}" % NS["p"]

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
px = lambda v: Emu(int(v * prs.slide_width / 1920))
blank = prs.slide_layouts[6]

import glob
for i in range(1, len(glob.glob(f"{renders}/slide*.png")) + 1):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(f"{renders}/slide{i}.png", 0, 0, prs.slide_width, prs.slide_height)
    if i == 6:
        # video sits exactly inside the bracket frame of the HTML slide: x 146, y 298, 804 x 603 px
        movie = slide.shapes.add_movie(video, px(146), px(298), px(804), px(603),
                                       poster_frame_image=poster, mime_type="video/mp4")
        spid = movie.shape_id
        timing = slide._element.find(P + "timing")
        if timing is not None:
            slide._element.remove(timing)
        slide._element.append(etree.fromstring(f'''
<p:timing xmlns:p="{NS['p']}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>
  <p:par><p:cTn id="2" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>
    <p:par><p:cTn id="3" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
      <p:par><p:cTn id="4" presetID="1" presetClass="mediacall" presetSubtype="0" fill="hold" nodeType="afterEffect"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>
        <p:cmd type="call" cmd="playFrom(0.0)"><p:cBhvr><p:cTn id="5" dur="20250" fill="hold"/><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cBhvr></p:cmd>
      </p:childTnLst></p:cTn></p:par>
    </p:childTnLst></p:cTn></p:par>
  </p:childTnLst></p:cTn></p:par>
  <p:video><p:cMediaNode vol="80000" mute="1"><p:cTn id="6" repeatCount="indefinite" fill="hold" display="0"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst></p:cTn><p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl></p:cMediaNode></p:video>
</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'''))
    # fade transition goes right after clrMapOvr (before timing)
    trans = etree.fromstring(f'<p:transition xmlns:p="{NS["p"]}" spd="med"><p:fade/></p:transition>')
    clr = slide._element.find(P + "clrMapOvr")
    if clr is not None:
        clr.addnext(trans)
    else:
        slide._element.find(P + "cSld").addnext(trans)

prs.core_properties.title = "Early Wildfire Detection"
prs.core_properties.author = "Kanatbay Alibek, Shaykhmetov Alimzhan"
prs.save(out)
print("saved", out)
